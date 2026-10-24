from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_hyperlink(paragraph, url, text, font_name='Times New Roman', font_size=10,
                   color=RGBColor(0x05, 0x63, 0xC1), bold=False, italic=False, underline=True):
    """
    Insert a real, clickable hyperlink run into a paragraph (python-docx has no
    built-in hyperlink API, so this builds the required w:hyperlink XML by hand).
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True
    )

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rPr.append(rFonts)

    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(font_size * 2)))
    rPr.append(sz)

    color_el = OxmlElement('w:color')
    color_el.set(qn('w:val'), str(color))
    rPr.append(color_el)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    if bold:
        rPr.append(OxmlElement('w:b'))

    if italic:
        rPr.append(OxmlElement('w:i'))

    new_run.append(rPr)

    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def set_cell_margins(cell, top=0, bottom=0, start=0, end=0):
    """
    Helper function to set cell margins tightly for the layout.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')

    for side, value in [('top', top), ('bottom', bottom), ('left', start), ('right', end)]:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(int(value * 1440)))  # 1440 twips = 1 inch
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)

    tcPr.append(tcMar)

def fix_table_widths(table, widths):
    """
    Word ignores python-docx cell widths unless the table uses a fixed layout and
    every grid column / cell carries the width, so set all of them explicitly.
    """
    tblPr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)

    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(sum(widths) * 1440)))
    tblW.set(qn('w:type'), 'dxa')

    for gridCol, w in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), widths):
        gridCol.set(qn('w:w'), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)

def create_resume():
    doc = Document()

    # --- 1. SETUP: Margins & Fonts ---
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)

    accent_color = RGBColor(46, 85, 160)

    # --- 2. HEADER SECTION ---
    header_table = doc.add_table(rows=1, cols=2)
    fix_table_widths(header_table, [4.9, 2.37])

    cell_h1 = header_table.cell(0, 0)
    name_run = cell_h1.paragraphs[0].add_run("Jamil Ahmad Rupak")
    name_run.bold = True
    name_run.font.size = Pt(26)

    title_run = cell_h1.add_paragraph().add_run("Software Engineer at Algoclan · Full-Stack Developer")
    title_run.font.size = Pt(13)
    title_run.italic = True

    cell_h2 = header_table.cell(0, 1)

    def contact_line(first=False):
        p = cell_h2.paragraphs[0] if first else cell_h2.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return p

    contact_line(first=True).add_run("Dhaka, Bangladesh").font.size = Pt(9)
    contact_line().add_run("+880 1771-823979").font.size = Pt(9)
    add_hyperlink(contact_line(), "mailto:jhrupok@gmail.com", "jhrupok@gmail.com", font_size=9)
    add_hyperlink(contact_line(), "https://linkedin.com/in/jamil-ahmad-rupak",
                  "linkedin.com/in/jamil-ahmad-rupak", font_size=9)
    add_hyperlink(contact_line(), "https://github.com/JamilAhmadRupak",
                  "github.com/JamilAhmadRupak", font_size=9)
    add_hyperlink(contact_line(), "https://jamilahmadrupak.vercel.app",
                  "jamilahmadrupak.vercel.app", font_size=9)

    # --- 3. MAIN LAYOUT (2 Columns) ---
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    layout_table = doc.add_table(rows=1, cols=2)
    fix_table_widths(layout_table, [5.0, 2.27])

    col1 = layout_table.cell(0, 0)
    set_cell_margins(col1, end=0.22)

    col2 = layout_table.cell(0, 1)

    def add_section_header(cell, text, first=False):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.color.rgb = accent_color
        run.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(0 if first else 10)
        p.paragraph_format.space_after = Pt(4)

    def add_bullet(cell, text):
        p = cell.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(text)
        return p

    def add_entry(cell, title, right, subtitle=None, space_before=6):
        """Bold title with right-aligned date/stack, optional italic subtitle underneath."""
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(4.7), alignment=2)  # 2 = RIGHT
        p.add_run(title).bold = True
        if right:
            r = p.add_run(f"\t{right}")
            r.italic = True
            r.font.size = Pt(9)
        if subtitle:
            s = p.add_run(f"\n{subtitle}")
            s.italic = True
            s.font.size = Pt(9)

    # === LEFT COLUMN ===
    add_section_header(col1, "Profile", first=True)
    summary = (
        "Full-stack software engineer building production inventory, order and pricing systems for "
        "e-commerce clients at Algoclan. Final-year CSE student at CUET; 1st place at the Mysoftheaven × "
        "CUET AI Hackathon, and top 10 in 3+ national hackathons. I like turning messy real-world "
        "workflows into reliable software."
    )
    p = col1.add_paragraph(summary)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_section_header(col1, "Experience")
    add_entry(col1, "Software Engineer, Algoclan", "Oct 2025 – Present",
              "Node.js, Express, Next.js, FastAPI, MongoDB, Docker", space_before=0)
    for text in [
        "Core contributor (250+ commits across backend and frontend) to a multi-tenant inventory and "
        "listing platform for an e-commerce reseller, integrated with two major resale-marketplace APIs.",
        "Built the RFID/barcode scanning workflow end to end: batch scanning, EPC-to-UPC conversion, "
        "reviewable scan drafts, fallback lookups and retries for unknown items, and commit-to-inventory.",
        "Built stock reconciliation against physical warehouse counts and applied ~19,400 reconciled "
        "inventory rows to production, with rollback files prepared in advance.",
        "Automated order ingestion and cross-marketplace delisting to prevent overselling; added "
        "multi-client onboarding and per-client attribution.",
        "Implemented four pricing strategies (competitive index, demand momentum, seasonality, market "
        "momentum) in a FastAPI repricing service, with scheduled bid/ask time-series collection and "
        "AI-generated price notifications.",
        "Added audit logging, rate limiting, API-key auth, Swagger docs and an AI weekly-analysis "
        "chatbot; Dockerized and set up CI/CD for another product's backend.",
    ]:
        add_bullet(col1, text)

    add_entry(col1, "AI R&D Industrial Trainee (Team Lead), Mysoftheaven (BD) Ltd.", "Sep 2026")
    add_bullet(col1, "Led a 3-person team building an AI Security Gateway against LLM prompt injection "
                     "and data leakage (OWASP LLM01/02): leakage-safe pipeline over 26k records, layered "
                     "regex/ML/DeBERTa detection, and an audit replacing an inflated 100% detection claim "
                     "with honest held-out metrics.")

    add_section_header(col1, "Personal Projects")
    add_entry(col1, "LTSR: Learning Through Spaced Repetition", "React, Django REST, Celery",
              space_before=0)
    add_bullet(col1, "Study app with SM-2 spaced-repetition flashcards, topic notes and recurring "
                     "tasks; Celery jobs send email reminders, with JWT auth and Chart.js progress charts.")

    add_entry(col1, "HealthTwin: AI-Powered Healthcare Platform", "Next.js, Django, PostgreSQL")
    add_bullet(col1, "Patient–doctor platform with symptom tracking, prescription assistance and "
                     "AI-generated patient insights, built around a \"digital twin\" of each "
                     "patient's health history.")

    add_entry(col1, "ModelMate: AI Community Platform", "Django, PostgreSQL")
    add_bullet(col1, "Community web app where AI users share knowledge and discuss models, with "
                     "authentication and a REST API.")

    add_section_header(col1, "Publication")
    pub = col1.add_paragraph()
    pub.add_run("Momentum: An E-Court Service and Lawyer Help Assistant Mobile Application").bold = True
    venue = pub.add_run("\nPresented at 16th ICCCNT 2025, IIT Indore (with IEEE EPS & AICTE) · Jul 2025")
    venue.italic = True
    venue.font.size = Pt(9)

    # === RIGHT COLUMN ===
    def add_tight_bullet(cell, text):
        p = cell.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"•\t{text}")

    add_section_header(col2, "Skills", first=True)
    skills_list = [
        ("Languages", "JavaScript, Python, Java, C++, C# (coursework)"),
        ("Backend", "Node.js, Express, FastAPI, Django REST Framework"),
        ("Frontend", "React, Next.js, Tailwind CSS, shadcn/ui"),
        ("Databases", "MongoDB, PostgreSQL, MySQL"),
        ("AI", "Vercel AI SDK, OpenAI & Gemini APIs"),
        ("DevOps & Tools", "Docker, CI/CD, Git, Swagger/OpenAPI, cron jobs, Celery"),
        ("Fundamentals", "Data Structures & Algorithms, OOP (C#), REST API design"),
    ]
    for cat, items in skills_list:
        p = col2.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f"{cat}: ").bold = True
        p.add_run(items)

    add_section_header(col2, "Education")
    e1 = col2.add_paragraph()
    e1.add_run("B.Sc. in Computer Science & Engineering").bold = True
    e1.add_run("\nChittagong University of Engineering & Technology (CUET)")
    e1.add_run("\nSession 2021–22 · Final year").italic = True

    e2 = col2.add_paragraph()
    e2.paragraph_format.space_before = Pt(6)
    e2.add_run("Higher Secondary Certificate").bold = True
    e2.add_run("\nBAF Shaheen College, Dhaka")
    e2.add_run("\nGPA 5.00 · 2021").italic = True

    e3 = col2.add_paragraph()
    e3.paragraph_format.space_before = Pt(6)
    e3.add_run("Python (Django) Training, 80 hrs").bold = True
    e3.add_run("\nEDGE Project, Bangladesh Computer Council, at CUET CSE")
    e3.add_run("\nSep – Dec 2024").italic = True

    add_section_header(col2, "Achievements")
    for item in [
        "1st Place, AI Hackathon by Mysoftheaven (BD) Ltd. & CUET CSE (2026)",
        "Top 10 in 3+ national hackathons, incl. Micro-Ops, CUET CSE Fest 2025",
        "Judge, BUILD@CUET Hackathon by MLH & Gemini (2026)",
        "~500 problems solved on Codeforces & LeetCode",
        "Former Secretary, BAF Shaheen College Math Club",
        "President's Scout Award; Shapla Cub Award",
    ]:
        add_tight_bullet(col2, item)

    doc.save('Jamil_Ahmad_Rupak.docx')


if __name__ == "__main__":
    create_resume()

